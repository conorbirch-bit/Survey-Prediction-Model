"""Standalone access workbook generated from the scheduler's existing decisions."""
from collections import Counter
from datetime import date, datetime
from io import BytesIO
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


def clean(value):
    if isinstance(value, (list, tuple)):
        return ', '.join(map(str, value))
    if value is None or pd.isna(value):
        return ''
    return value


def excel_date(value):
    parsed = pd.to_datetime(value, errors='coerce')
    return parsed.date() if not pd.isna(parsed) else None


def action_group(row):
    decision = row.get('Decision')
    if decision == 'RESOLVED':
        return 'Completed — no further action'
    if decision == 'CLIENT_ACCESS_REQUIRED':
        return 'Client access required'
    if decision in ('RETRY', 'RETRY_WITH_CONSTRAINT'):
        if clean(row.get('Booking Excluded')) == True:
            return 'Already booked in excluded week'
        if not str(row.get('Mapping Status', '')).startswith('OK'):
            return 'Replacement appointment check'
        return 'Routine retry'
    return 'Internal review / hold'


def table_sheet(wb, name, headers, rows, widths=None):
    ws = wb.create_sheet(name)
    ws.cell(1, 1, name).font = Font(name='Arial', size=18, bold=True, color='17365D')
    ws.cell(3, 1, f'{len(rows)} records')
    for col, header in enumerate(headers, 1):
        cell = ws.cell(7, col, header)
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = (widths or {}).get(col, 25)
    for index, row in enumerate(rows, 8):
        for col, value in enumerate(row, 1):
            value = clean(value)
            cell = ws.cell(index, col, value)
            # Surveyor prose and identifiers are literal text, including '=...'.
            if isinstance(value, str):
                cell.data_type = 's'
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if isinstance(value, (date, datetime)):
                cell.number_format = 'dd mmm yyyy'
        ws.row_dimensions[index].height = 94
    ws.row_dimensions[7].height = 42
    if rows:
        table = Table(displayName=name.replace(' ', ''), ref=f'A7:{get_column_letter(len(headers))}{7+len(rows)}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = 'C8'
    ws.sheet_view.showGridLines = False
    return ws


def build_access_report(plan, source_name='', generated_on=None):
    """Return XLSX bytes. Booking selection never removes client-help cases."""
    frame = getattr(plan, 'report_decisions', pd.DataFrame())
    if frame.empty:
        frame = plan.decisions
    records = frame.to_dict('records')
    records.sort(key=lambda r: (action_group(r), str(clean(r.get('Customer Reference'))), str(r.get('Work Order Number', ''))))
    visits = sorted({key for row in records for key in row if re.fullmatch(r'Visit \d+ Date', key)}, key=lambda key: int(key.split()[1]))
    visits = sorted(set(visits + ['Visit 1 Date', 'Visit 2 Date']), key=lambda key: int(key.split()[1]))
    clients = [r for r in records if r.get('Decision') == 'CLIENT_ACCESS_REQUIRED']
    wb = Workbook()
    summary = wb.active
    summary.title = 'Summary'
    summary.append(['Cannot Completes Access Report'])
    summary.append(['Generated', generated_on or date.today()])
    summary['B2'].number_format = 'dd mmm yyyy'
    summary.append(['Source', source_name])
    summary.append(['Register records (one per Work Order)', len(records)])
    summary.append(['Action group', 'Count'])
    counts = Counter(action_group(r) for r in records)
    for name, count in sorted(counts.items()):
        summary.append([name, count])
    summary.append([])
    summary.append(['Client access issue', 'Count'])
    for name, count in sorted(Counter(str(clean(r.get('Reason Category'))).replace('_', ' ').title() for r in clients).items()):
        summary.append([name, count])
    summary.append([])
    for note in [
        'Client help: two customer failures, or a clear access barrier after one visit. Metro failures do not count towards the customer limit.',
        'Any linked Completed service appointment resolves the Work Order and removes it from the client tab.',
        'Booking-week exclusions affect scheduling only; they do not hide client access issues.',
        'Blank failure narratives are assessed as No answer at door; original descriptions remain blank.',
        'Visit columns contain recorded failed-appointment Actual Start dates, in chronological order. Missing dates remain blank.',
        'The register covers the uploaded failure-detail Work Orders plus completed Work Orders found in the appointment mapping. Mapping-only records may lack building names, references and failure counts; blanks mean unavailable, not zero.',
        'Client response fields are blank in each new export. Keep the returned client workbook separately.',
    ] + list(plan.warnings):
        summary.append([note])
    summary.column_dimensions['A'].width = 110
    summary.column_dimensions['B'].width = 38
    for row in summary:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        summary.row_dimensions[row[0].row].height = 44
    summary['A1'].font = Font(name='Arial', size=18, bold=True, color='17365D')
    summary.freeze_panes = 'A6'
    headers = ['Customer Reference', 'Building / address', 'Postcode', 'Access issue', 'Why help is needed', 'Help requested', 'Recorded customer failures', 'Last failed visit', 'Client response / access instructions', 'Site contact', 'Agreed access date', 'Response status', 'Surveyor’s original description'] + [v.replace('Date', 'date') for v in visits]
    client_rows = [[r.get('Customer Reference'), r.get('Building Name'), r.get('Postcode'), str(clean(r.get('Reason Category'))).replace('_',' ').title(), r.get('Decision Reason'), r.get('Recommended Client Action'), r.get('Customer Failure Count'), excel_date(r.get('Previous Visit')), '', '', None, 'Awaiting response', r.get('Surveyor Original Description')] + [excel_date(r.get(v)) for v in visits] for r in clients]
    ws = table_sheet(wb, 'Client Access Required', headers, client_rows, {1:19,2:34,3:14,4:29,5:51,6:54,7:15,8:17,9:47,10:29,11:19,12:26,13:65, **{i:19 for i in range(14,len(headers)+1)}})
    validation = DataValidation(type='list', formula1='"Awaiting response,Contact provided,Access arranged,Further clarification needed,Unable to arrange"', allow_blank=True)
    ws.add_data_validation(validation)
    if clients:
        validation.add(f'L8:L{7+len(clients)}')
    for row in ws.iter_rows(min_row=8, max_row=7+len(clients), min_col=9, max_col=12):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='FFF2CC')
        row[2].number_format = 'dd mmm yyyy'
    fields = ['Customer Reference','Building Name','Postcode','Work Order Number','Decision','Reason Category','Decision Reason','Recommended Client Action','Customer Failure Count','Metro Failure Count','Previous Visit','Surveyor Original Description','Primary Service Appointment: Reason Description','Cancelation Reason Description','Reason Not Complete','Failure Reason','Mapping Status','Replacement Service Appointment ID','Old Service Appointment ID','Completed Service Appointment IDs','Linked Appointment Statuses','Booking Excluded','Retry Eligible','Required Weekdays','Forbidden Weekday','Preferred Retry Period','Record Coverage','Building ID'] + visits
    rows = [[action_group(r)] + [excel_date(r.get(k)) if k == 'Previous Visit' or k in visits else r.get(k) for k in fields] for r in records]
    table_sheet(wb, 'All Cannot Completes', ['Action group']+fields, rows, {7:50,8:50,13:65,14:65,15:65})
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
